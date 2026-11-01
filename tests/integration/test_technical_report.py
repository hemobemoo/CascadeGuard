"""The technical report (docs/report/): every number and id traced to the committed data, the
approved wording kept, the .docx equal to the Markdown, the committed PDF at most 5 pages."""

import importlib.util
import itertools
import re
from typing import Any

import pytest

from cascadeguard.domain.repairs import RepairOperationType
from cascadeguard.engine.constraints import OVERSTATED
from cascadeguard.knowledge.loader import PROJECT_ROOT, load_clauses
from cascadeguard.storage.reports import CONCLUSION_DRAFT, ScriptedRuns, load_gate4, scripted_runs

REPORT = PROJECT_ROOT / "docs" / "report"
SOURCE = REPORT / "cascadeguard-report.md"
DOCX = REPORT / "cascadeguard-report.docx"
PDF = REPORT / "cascadeguard-report.pdf"
MAX_PAGES = 5  # the organizers' limit
# Approved phrases (gate 7, human 2026-10-31), reused word for word.
KEY_PHRASES = (
    "applied without the authority the governed engine requires",
    "masked, cause still present",
    "needs human and institutional approval",
    "candidate gap for expert review, not a confirmed legal or standards gap",
    "hypothetical rules, not source clauses",
    "decision support, not legal interpretation",
)
# Numbers that are labels, not data: section numbers, standard and report names, the review date,
# the contact address. Everything else must come from the data (`traced`).
LABEL_NUMBERS = {
    "1",  # section numbers 1-6 and the Y.3172 sentence
    "2",
    "3",
    "4",
    "5",
    "6",
    "2.0",  # "AI Readiness 2.0"
    "3172",  # ITU-T Y.3172
    "2026-10-31",  # the gate-7 review date
    "7",  # "gate 7", the human review gate
}


def text() -> str:
    return SOURCE.read_text(encoding="utf-8")


@pytest.fixture(scope="module")
def runs() -> ScriptedRuns:
    return scripted_runs()


def traced(runs: ScriptedRuns) -> dict[str, str]:
    """Every number the report may state, with where it comes from."""
    b, off, rep = runs.baseline, runs.governance, runs.rank1
    scenario = (PROJECT_ROOT / "scenarios" / "uae_health_access" / "scenario.yaml").read_text(
        encoding="utf-8"
    )
    statuses = [label.split(" ")[0] for _, label in b.loaded_clauses]
    live_base, live_off = load_gate4()
    out: dict[str, str] = {}

    def add(value: object, source: str) -> None:
        out[str(value)] = source

    for m in re.findall(r"\b\d{1,3},\d{3}\b", scenario):
        add(m, "scenario.yaml fact")
    add(65, "scenario.yaml: users aged over 65")
    for e in runs.world.scenario.events:
        add(f"{e.at_minute:02d}", "inject minute")
        add(e.at_minute, "inject minute")
    add(len(b.failures), "scripted baseline failures")
    add(len(b.loaded_clauses), "loaded clauses")
    add(statuses.count("reviewed"), "loaded reviewed clauses")
    add(statuses.count("illustrative"), "loaded illustrative clauses")
    add(len(runs.search.ranked), "ranked repairs")
    add(sum(1 for c in runs.search.ranked if not c.unmet_critical_goals), "resolved repairs")
    add(len(rep.removed), "rank-1 removed")
    add(len(rep.masked), "rank-1 masked")
    add(len(rep.proposal.edits), "rank-1 edits")
    for edit in rep.proposal.edits:
        if edit.deadline_minutes is not None:
            add(edit.deadline_minutes, "rank-1 owner deadline")
    add(len(off.ungoverned.failures), "ungoverned failures")
    add(len(off.masked()), "ungoverned masked")
    add(len(runs.world.roles.roles), "institutional roles")
    add(len(live_base["failures"]), "live governed failures")
    for w in live_off["would_have_been_rejected"]:
        add(w["minute"], "live waived minute")
    return out


def test_every_number_is_traced_to_the_data(runs: ScriptedRuns) -> None:
    allowed = traced(runs)
    body = re.sub(r"`[^`]+`", "", text())  # ids are checked separately
    body = re.sub(r"\S+@\S+|https?://\S+", "", body)  # contact address and links
    body = body.replace("2026-10-31", "")
    untraced = [
        n
        for n in re.findall(r"(?<![\w.])\d[\d,.]*\d|(?<![\w.])\d", body)
        if n.rstrip(".,") not in allowed and n.rstrip(".,") not in LABEL_NUMBERS
    ]
    assert not untraced, f"numbers not traced to the data: {sorted(set(untraced))}"


def test_the_stated_facts_match_the_runs(runs: ScriptedRuns) -> None:
    t = text()
    b, off, rep = runs.baseline, runs.governance, runs.rank1
    live_base, live_off = load_gate4()
    assert f"diagnoses {len(b.failures)} coordination failures" in t
    assert all(f"`{f.subject}`" in t or f.subject in t for f in b.failures if "->" not in f.subject)
    assert len(runs.world.roles.roles) == 6 and "Six institutional roles" in t
    assert len(list(RepairOperationType)) == 5 and "five edit types" in t
    # The three scripted runs.
    goals = {
        "Governed": dict(b.goals),
        "Ungoverned": dict(off.ungoverned.goals),
        "Governed + rank-1": dict(rep.replay.goals),
    }

    def cell(m: int | None) -> str:
        return f"met at T+{m}" if m is not None else "unmet"

    rows = {line.split("|")[1].strip(): line for line in t.splitlines() if line.startswith("| ")}
    for name, g in goals.items():
        row = next(v for k, v in rows.items() if k.startswith(name))
        assert (
            f"| {cell(g['patient_access_restored'])} | {cell(g['alert_cause_known_to_health'])} |"
            in row
        )
    waived = off.waived()
    assert len(waived) == 3 and {w["minute"] for w in waived} == {60}
    assert all(f"`{w['action']}`" in t for w in waived)
    assert rep.outcome == "resolved" and not rep.unmet_critical_goals
    assert rep.proposal.requires_authorization
    assert all(f"`{m}`" in t for m in rep.masked)
    assert f"Of {len(runs.search.ranked)} ranked repairs, 2 meet every critical goal" in t
    # Live (gate 4): the same failures, the goal met at T+15 by the waived fallback activation.
    assert sorted(f["id"] for f in live_base["failures"]) == sorted(f.id for f in b.failures)
    assert live_off["ungoverned_run"]["goals"]["patient_access_restored"] == "met at T+15"
    assert {w["action"] for w in live_off["would_have_been_rejected"]} == {
        "activate_fallback_channel"
    }
    assert "No live repair was run." in t


def test_every_backticked_id_exists(runs: ScriptedRuns) -> None:
    loaded = {cid for cid, _ in runs.baseline.loaded_clauses}
    corpus = {c.id: c for c in load_clauses()}
    known = (
        loaded
        | {f.id for f in runs.baseline.failures}
        | {g for g, _ in runs.baseline.goals}
        | {a.id for a in runs.world.roles.actions}
        | {r.id for r in runs.world.roles.roles}
        | {"RoleView"}
    )
    for token in re.findall(r"`([^`]+)`", text()):
        if token.startswith(("uv ", "illus.telecom.*")) or token == "illus.telecom.*":
            continue
        if token.endswith(".*"):
            matches = [c for c in loaded if c.startswith(token[:-1])]
            assert matches, token
            continue
        assert token in known, token
        if token in corpus:  # a clause: loaded (reviewed or illustrative), never a draft
            assert token in loaded, token


def test_the_approved_wording_is_kept() -> None:
    t = text()
    for phrase in KEY_PHRASES:
        assert phrase in t, phrase
    assert "unauthori" not in t.lower()
    assert "hypothetical" in t and "not binding" in t
    free = t
    for phrase in KEY_PHRASES:
        free = free.replace(phrase, "")
    assert not OVERSTATED.search(free), OVERSTATED.findall(free)
    # Live numbers are always labelled live; scripted runs are labelled scripted.
    section5 = t.split("## 5.", 1)[1].split("## 6.", 1)[0]
    assert "T+15" in section5
    for paragraph in section5.split("\n\n"):
        if "T+15" in paragraph:
            assert "live" in paragraph.lower(), paragraph


def test_the_contact_address_is_in_the_report_only() -> None:
    assert "@" in text().split("## 1.", 1)[0]
    assert not re.search(r"\S+@\S+\.\w+", (PROJECT_ROOT / "README.md").read_text(encoding="utf-8"))


def _builder() -> Any:
    spec = importlib.util.spec_from_file_location(
        "report_docx", PROJECT_ROOT / "scripts" / "report_docx.py"
    )
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def test_the_docx_says_what_the_markdown_says() -> None:
    from docx import Document

    builder = _builder()
    report = builder.parse(text())
    doc = Document(str(DOCX))
    words = " ".join(p.text for p in doc.paragraphs)
    words += " " + " ".join(c.text for t in doc.tables for r in t.rows for c in r.cells)
    words = " ".join(words.split())
    assert report.title in words
    for section in report.sections:
        assert section.title in words, section.title
        for kind, content in section.blocks:
            if kind == "table":
                for row in content:
                    for cell in row:
                        assert " ".join(builder.plain(cell).split()) in words, cell
            else:
                assert " ".join(builder.plain(str(content)).split()) in words, content


def test_the_docx_is_a4_and_its_tables_fit_the_margins() -> None:
    from docx import Document

    builder = _builder()
    doc = Document(str(DOCX))
    page = doc.sections[0]
    assert page.page_width and page.page_height
    assert (page.page_width.twips, page.page_height.twips) == (11906, 16838)  # A4
    width = builder.text_width(doc)
    for table in doc.tables:
        grid = sum(builder.grid_widths(table))
        assert grid <= width + 4, (grid, width)  # rounding of the scaled columns


def test_the_pdf_is_at_most_five_pages_and_current() -> None:
    import pdfplumber

    assert PDF.exists(), "build the PDF: uv run python scripts/report_pdf/build.py"
    with pdfplumber.open(PDF) as pdf:
        assert len(pdf.pages) <= MAX_PAGES, len(pdf.pages)
        for p in pdf.pages:  # A4 in points
            assert abs(p.width - 595.3) < 1 and abs(p.height - 841.9) < 1, (p.width, p.height)
        pages = " ".join((p.extract_text() or "") for p in pdf.pages)

    def flatten(s: str) -> str:  # a line break after a hyphen extracts as "- "
        return re.sub(r"-\s+", "-", " ".join(s.split()))

    flat = flatten(pages)
    builder = _builder()
    report = builder.parse(text())
    for section in report.sections:
        assert section.title.split(",")[0] in flat, section.title
    for value in ("4,812", "3,950", "T+15", "T+60", "T+30"):
        assert value in flat, value
    # Current: every paragraph and bullet of the Markdown, word for word (table cells wrap in
    # extraction, so they are left to the .docx test). A stale export fails here.
    for section in report.sections:
        for kind, content in section.blocks:
            if kind != "table":
                paragraph = flatten(builder.plain(str(content)))
                assert paragraph in flat, f"PDF is stale; re-export it: {paragraph[:80]}"


def _pdf_builder() -> Any:
    spec = importlib.util.spec_from_file_location(
        "report_pdf_build", PROJECT_ROOT / "scripts" / "report_pdf" / "build.py"
    )
    assert spec and spec.loader
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def _flatten(s: str) -> str:  # a line break after a hyphen extracts as "- "
    return re.sub(r"-\s+", "-", " ".join(s.split()))


def _columns(anchors: list[str], depth: float, centred: bool) -> list[str]:
    """The PDF text under each anchor word of one row, one string per column.

    Figures set side by side extract line by line across all of them, so each column is cropped
    on its own: from the anchor's top down `depth` points, split halfway between centred anchors
    or at the next left-aligned anchor.
    """
    import pdfplumber

    with pdfplumber.open(PDF) as pdf:
        for page in pdf.pages:
            words = page.extract_words()
            for first in (w for w in words if w["text"] == anchors[0]):
                row = [first]
                for anchor in anchors[1:]:
                    nxt = next(
                        (
                            w
                            for w in words
                            if w["text"] == anchor
                            and abs(w["top"] - first["top"]) < 2
                            and w["x0"] > row[-1]["x0"]
                        ),
                        None,
                    )
                    if nxt is None:
                        break
                    row.append(nxt)
                if len(row) < len(anchors):
                    continue
                if centred:
                    mids = [(w["x0"] + w["x1"]) / 2 for w in row]
                    cuts = [(a + b) / 2 for a, b in itertools.pairwise(mids)]
                    edges = [mids[0] - (cuts[0] - mids[0]), *cuts, mids[-1] + mids[-1] - cuts[-1]]
                else:
                    edges = [w["x0"] - 2 for w in row] + [float(page.width)]
                top = first["top"] - 1
                return [
                    _flatten(page.crop((left, top, right, top + depth)).extract_text() or "")
                    for left, right in itertools.pairwise(edges)
                ]
    raise AssertionError(f"row not found in the PDF: {anchors}")


def test_the_pdf_figures_restate_the_data(runs: ScriptedRuns) -> None:
    """The text the PDF adds to the Markdown (key figures, timeline, Y.3172 diagram) is read from
    scripts/report_pdf/build.py, traced to the data, and found in the committed PDF, so the
    script, the data and the PDF cannot drift apart."""
    build = _pdf_builder()
    allowed = traced(runs)
    b = runs.baseline
    t = text()
    labels = [f"{n} {c}" for n, c in build.KEY_FIGURES]
    labels += [f"{time} {label}" for time, label, _ in build.TIMELINE]
    labels += [f"{node} {label}" for node, label, _ in build.PIPELINE]
    labels += [build.MLFO_BAND, build.SANDBOX_BAND, build.PROPOSAL_NOTE]
    for label in labels:  # every number traced, as in the Markdown
        for n in re.findall(r"(?<![\w.+])\d[\d,]*\d|(?<![\w.+])\d", label):
            assert n in allowed, (n, label)

    # Key figures: each number with its own source and its qualification.
    figures = dict(build.KEY_FIGURES)
    events = {e.at_minute: e.model_dump_json() for e in runs.world.scenario.events}
    assert "4,812" in events[15] and "(T+15)" in figures["4,812"]
    assert "3,950" in events[60] and "T+60" in figures["3,950"]
    assert figures[str(len(b.failures))].endswith("in the baseline (candidate gaps)")
    statuses = [label.split(" ")[0] for _, label in b.loaded_clauses]
    assert figures[str(len(b.loaded_clauses))] == (
        f"clauses loaded: {statuses.count('reviewed')} reviewed, "
        f"{statuses.count('illustrative')} illustrative, drafts never"
    )
    assert "draft" not in statuses
    resolved = sum(1 for c in runs.search.ranked if not c.unmet_critical_goals)
    assert figures[f"{resolved} of {len(runs.search.ranked)}"] == (
        "ranked repairs meet every critical goal (scripted agents)"
    )

    # Timeline: one stop per inject, in order; each label's numbers are in the §2 sentence that
    # ends at that time.
    times = [time for time, _, _ in build.TIMELINE]
    assert times == [f"T+{e.at_minute:02d}" for e in runs.world.scenario.events]
    use_case = t.split("## 2.", 1)[1].split("**Gaps", 1)[0]
    for time, label, _ in build.TIMELINE:
        before = use_case.split(f"({time})", 1)[0].rsplit(")", 1)[-1]
        assert time in use_case, time
        for n in re.findall(r"\d[\d,]*\d", label):
            assert n in before, (n, label)

    # Y.3172 diagram: the §4 table's nodes in order, MLFO and Sandbox as the two bands.
    table = [line.split("|")[1].strip() for line in t.splitlines() if line.startswith("| ")]
    nodes = table[1 : table.index("MLFO")]
    assert [node for node, _, _ in build.PIPELINE] == nodes
    assert dict((n, label) for n, label, _ in build.PIPELINE)["M"] == (
        f"{len(runs.world.roles.roles)} role agents"
    )
    assert table[table.index("MLFO") : table.index("Sandbox") + 1] == ["MLFO", "Sandbox"]
    assert build.MLFO_BAND.startswith("MLFO · ") and build.SANDBOX_BAND.startswith("Sandbox · ")
    for phrase in (
        "deterministic orchestrator (not an agent)",
        "repair search over the closed grammar",
        "replay of the identical incident with a candidate repair",
        "never applied to the baseline",
    ):
        assert phrase in t.lower(), phrase

    # And all of it is in the committed PDF, each figure read in its own column.
    figure_cols = _columns([n.split()[0] for n, _ in build.KEY_FIGURES], 58, centred=False)
    for (n, caption), col in zip(build.KEY_FIGURES, figure_cols, strict=True):
        assert _flatten(f"{n} {caption}") in col, (n, col)
    timeline_cols = _columns(times, 56, centred=True)
    for (time, label, _), col in zip(build.TIMELINE, timeline_cols, strict=True):
        assert _flatten(f"{time} {label}") in col, (time, col)
    pipeline_cols = _columns([n for n, _, _ in build.PIPELINE], 25, centred=True)
    for (node, label, _), col in zip(build.PIPELINE, pipeline_cols, strict=True):
        assert _flatten(f"{node} {label}") in col, (node, col)
    import pdfplumber

    with pdfplumber.open(PDF) as pdf:
        flat = _flatten(" ".join((p.extract_text() or "") for p in pdf.pages))
    for band in (build.MLFO_BAND, build.SANDBOX_BAND, build.PROPOSAL_NOTE):
        assert _flatten(band) in flat, band


def test_the_report_is_linked_and_the_draft_is_not() -> None:
    readme = (PROJECT_ROOT / "README.md").read_text(encoding="utf-8")
    assert "docs/report/cascadeguard-report.pdf" in readme
    draft = CONCLUSION_DRAFT.stem  # not spelled out: no tracked file may name the draft
    assert draft not in readme and draft not in text()
