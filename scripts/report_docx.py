# pyright: reportPrivateUsage=false, reportUnknownMemberType=false
# pyright: reportAttributeAccessIssue=false, reportUnknownVariableType=false
# pyright: reportUnknownArgumentType=false
# The builder edits the template's XML through python-docx internals (_p, _tbl, _r):
# python-docx has no public API to move or remove paragraphs and tables.
"""Build the technical report .docx from its Markdown source, on the organizers' template.

The template's title line, team lines and six numbered section headings are kept and rewritten
in place, with their styles; its example body text is removed; its mapping table (style "Grid
Table Light") is refilled with the report's Y.3172 mapping. Other Markdown tables use the same
table style. Only the Markdown subset the report uses is supported: `#`/`##` headings, "Label:
value" header lines, paragraphs, "- " bullets, pipe tables, **bold** and `code`.

    uv run python scripts/report_docx.py [--template PATH]

The organizers' template is not in the repository: it is read from the folder that contains the
checkout (pass --template otherwise), so a clone can check the committed .docx but not rebuild it.
Writes docs/report/cascadeguard-report.docx. The PDF is built from the Markdown separately, with
Chromium (scripts/report_pdf/build.py). python-docx is a dev-only dependency (human gate 3).
"""

import argparse
import copy
import re
import sys
from dataclasses import dataclass, field
from datetime import UTC, datetime
from pathlib import Path

from docx import Document
from docx.document import Document as DocumentType
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Mm, Pt
from docx.table import Table
from docx.text.paragraph import Paragraph

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "docs" / "report" / "cascadeguard-report.md"
OUTPUT = ROOT / "docs" / "report" / "cascadeguard-report.docx"
TEMPLATE = ROOT.parent / "AI-for-Good-Sandbox-Hackathon-submission-template.docx"
FONT = "Times New Roman"
CODE_FONT = "Consolas"
BODY_PT = 10.5
TABLE_PT = 8.5
FIXED_TIME = datetime(2026, 10, 31, tzinfo=UTC)  # deterministic document properties
# The template's six section headings, in order (numbered "List Paragraph" paragraphs).
TEMPLATE_HEADINGS = (
    "Introduction",
    "Description of the use case",
    "Use case requirement",
    "Mapping with ITU-T Y.3172",
    "Error scenarios",
    "Policy gaps and standards gaps",
)
TEAM_LABELS = ("Team leader:", "Team members:", "Contact details:", "Github link:")


@dataclass
class Section:
    title: str
    blocks: list[tuple[str, object]] = field(default_factory=list)  # (kind, content)


@dataclass
class Report:
    title: str
    team: dict[str, str]
    sections: list[Section]


def parse(markdown: str) -> Report:
    title = ""
    team: dict[str, str] = {}
    sections: list[Section] = []
    lines = markdown.splitlines()
    i = 0
    while i < len(lines):
        line = lines[i].rstrip()
        if line.startswith("# "):
            title = line[2:].strip()
        elif line.startswith("## "):
            sections.append(Section(re.sub(r"^\d+\.\s*", "", line[3:].strip())))
        elif not sections and any(line.startswith(label) for label in TEAM_LABELS):
            label, _, value = line.partition(":")
            team[label + ":"] = value.strip()
        elif line.startswith("|") and sections:
            rows: list[list[str]] = []
            while i < len(lines) and lines[i].startswith("|"):
                cells = [c.strip() for c in lines[i].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", c) for c in cells):
                    rows.append(cells)
                i += 1
            sections[-1].blocks.append(("table", rows))
            continue
        elif line.startswith("- ") and sections:
            sections[-1].blocks.append(("bullet", line[2:].strip()))
        elif line.strip() and sections:
            sections[-1].blocks.append(("para", line.strip()))
        i += 1
    return Report(title, team, sections)


def plain(text: str) -> str:
    """The text a reader sees: Markdown emphasis and code marks removed."""
    return re.sub(r"\*\*(.+?)\*\*|`([^`]+)`", lambda m: m.group(1) or m.group(2), text)


def _runs(paragraph: Paragraph, text: str, size: float) -> None:
    for part in re.split(r"(\*\*.+?\*\*|`[^`]+`)", text):
        if not part:
            continue
        if part.startswith("**"):
            run = paragraph.add_run(part[2:-2])
            run.bold = True
            run.font.name = FONT
        elif part.startswith("`"):
            run = paragraph.add_run(part[1:-1])
            run.font.name = CODE_FONT
            size_here = size - 1
            run.font.size = Pt(size_here)
            continue
        else:
            run = paragraph.add_run(part)
            run.font.name = FONT
        run.font.size = Pt(size)


def _new_paragraph(doc: DocumentType, after: object, text: str, bullet: bool) -> object:
    p = doc.add_paragraph()
    fmt = p.paragraph_format
    fmt.space_after = Pt(3)
    fmt.space_before = Pt(0)
    if bullet:
        fmt.left_indent = Cm(0.5)
        fmt.first_line_indent = Cm(-0.3)
        text = "• " + text
    _runs(p, text, BODY_PT)
    after.addnext(p._p)  # type: ignore[attr-defined]
    return p._p


def _fill_table(table: Table, rows: list[list[str]]) -> None:
    while len(table.rows) > 1:
        table._tbl.remove(table.rows[-1]._tr)
    template_row = copy.deepcopy(table.rows[0]._tr)
    table._tbl.remove(table.rows[0]._tr)
    for cells in rows:
        tr = copy.deepcopy(template_row)
        table._tbl.append(tr)
        row = table.rows[-1]
        for cell, text in zip(row.cells, cells, strict=True):
            for p in cell.paragraphs[1:]:
                p._p.getparent().remove(p._p)
            p = cell.paragraphs[0]
            for r in list(p.runs):
                r._r.getparent().remove(r._r)
            p.paragraph_format.space_after = Pt(0)
            _runs(p, text, TABLE_PT)


def text_width(doc: DocumentType) -> int:
    """The first section's text width between the margins, in twips."""
    page = doc.sections[0]
    assert page.page_width and page.left_margin is not None and page.right_margin is not None
    return Emu(page.page_width - page.left_margin - page.right_margin).twips


def grid_widths(table: Table) -> list[int]:
    """A table's grid column widths, in twips."""
    return [int(str(col.get(qn("w:w")))) for col in table._tbl.tblGrid.findall(qn("w:gridCol"))]


def _fit_table(table: Table, width: int) -> None:
    """Scale a table's fixed widths (table, grid columns, cells; in twips) to the text width."""
    tbl = table._tbl
    total = sum(grid_widths(table))
    if not total:
        return
    scale = width / total
    for el in [*tbl.tblGrid.findall(qn("w:gridCol")), *tbl.iter(qn("w:tcW"))]:
        if el.get(qn("w:type"), "dxa") == "dxa" and el.get(qn("w:w")):
            el.set(qn("w:w"), str(round(int(str(el.get(qn("w:w")))) * scale)))
    tbl_w = tbl.tblPr.find(qn("w:tblW"))
    if tbl_w is not None and tbl_w.get(qn("w:type")) == "dxa":
        tbl_w.set(qn("w:w"), str(width))


def _new_table(doc: DocumentType, after: object, rows: list[list[str]], style: str) -> object:
    table = doc.add_table(rows=1, cols=len(rows[0]))
    table.style = style
    _fill_table(table, rows)
    after.addnext(table._tbl)
    return table._tbl


def build(markdown: str, template: Path) -> DocumentType:
    report = parse(markdown)
    doc = Document(str(template))
    body = doc.paragraphs
    # Title: keep the first run's formatting, replace the text.
    title = body[0]
    for r in title.runs[1:]:
        r._r.getparent().remove(r._r)
    title.runs[0].text = report.title
    # Team lines: the template has "Team leader:", "Team members:\nContact details:", "Github".
    for p in body[1:4]:
        text = p.text
        for r in list(p.runs):
            r._r.getparent().remove(r._r)
        lines = [label for label in TEAM_LABELS if label in text]
        if len(lines) > 1:  # a justified line ending in a break is stretched: align left
            p.paragraph_format.alignment = None
        for n, label in enumerate(lines):
            if n:
                p.add_run().add_break()
            head = p.add_run(label + " ")
            head.bold = True
            head.font.name = FONT
            value = p.add_run(report.team.get(label, ""))
            value.font.name = FONT
    # Section headings, found by their template text, rewritten with the report's titles.
    headings: list[Paragraph] = []
    for expected in TEMPLATE_HEADINGS:
        match = next(p for p in body if p.text.strip().startswith(expected))
        headings.append(match)
    if len(report.sections) != len(headings):
        raise ValueError(f"the report needs {len(headings)} sections, got {len(report.sections)}")
    keep = {id(p._p) for p in (title, *body[1:4], *headings)}
    template_table = doc.tables[0]
    template_table._tbl.getparent().remove(template_table._tbl)
    for p in body:
        if id(p._p) not in keep and p._p.getparent() is not None:
            p._p.getparent().remove(p._p)
    for heading, section in zip(headings, report.sections, strict=True):
        for r in heading.runs[1:]:
            r._r.getparent().remove(r._r)
        heading.runs[0].text = section.title
        heading.paragraph_format.space_before = Pt(6)
        heading.paragraph_format.space_after = Pt(2)
        anchor = heading._p
        mapping_used = False
        for kind, content in section.blocks:
            if kind == "table":
                rows = content  # type: ignore[assignment]
                if section.title.startswith("Mapping") and not mapping_used:
                    _fill_table(template_table, rows)  # type: ignore[arg-type]
                    anchor.addnext(template_table._tbl)
                    anchor = template_table._tbl
                    mapping_used = True
                else:
                    style = template_table.style.name if template_table.style else None
                    anchor = _new_table(doc, anchor, rows, style)  # type: ignore[arg-type]
            else:
                anchor = _new_paragraph(doc, anchor, str(content), bullet=kind == "bullet")
    for page in doc.sections:  # the template is US Letter; the report is A4, margins kept
        page.page_width = Mm(210)
        page.page_height = Mm(297)
    for table in doc.tables:
        _fit_table(table, text_width(doc))
    props = doc.core_properties
    props.title = report.title
    props.author = "GM of AI"
    props.last_modified_by = "GM of AI"
    props.created = FIXED_TIME
    props.modified = FIXED_TIME
    props.revision = 1
    return doc


def main() -> int:
    parser = argparse.ArgumentParser()
    parser.add_argument("--template", type=Path, default=TEMPLATE)
    args = parser.parse_args()
    if not args.template.exists():
        print(f"template not found: {args.template}")
        return 2
    doc = build(SOURCE.read_text(encoding="utf-8"), args.template)
    doc.save(str(OUTPUT))
    print(f"wrote {OUTPUT.relative_to(ROOT).as_posix()}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
